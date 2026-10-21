"""რეპორტის ექსპორტი Excel-ში (თითო სექცია — ცალკე ფურცელი, რიცხვები რიცხვებად რჩება)."""
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.i18n import t
from app.services.reports.data import Report

FORMATS = {
    "qty": "#,##0.####",
    "price": '#,##0.00## "₾"',
    "num": "#,##0.##",
    "money": '#,##0.00 "₾"',
    "percent": '0.0"%"',
    "date": "DD.MM.YYYY",
}
HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")
LOW_FONT = Font(color="C62828", bold=True)


def _sheet_name(title: str, used: set) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", title)[:28].strip() or "Sheet"
    base, n = name, 2
    while name in used:
        name = f"{base[:25]} {n}"
        n += 1
    used.add(name)
    return name


def to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    for sec in report.sections:
        ws = wb.create_sheet(_sheet_name(sec.title, used))
        ws.append([report.title])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"{sec.title} · {report.period}"])
        ws["A2"].font = Font(italic=True, color="555555")
        ws.append([])
        header_row = 4
        ws.append([c.label for c in sec.columns])
        for i, _ in enumerate(sec.columns, start=1):
            cell = ws.cell(row=header_row, column=i)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        widths = [len(c.label) + 2 for c in sec.columns]
        rows = list(sec.rows)
        flags = list(sec.row_flags) + [""] * (len(rows) - len(sec.row_flags))
        if sec.totals:
            rows.append(sec.totals)
            flags.append("total")
        for row, flag in zip(rows, flags):
            ws.append([row.get(c.key) for c in sec.columns])
            r = ws.max_row
            for i, c in enumerate(sec.columns, start=1):
                cell = ws.cell(row=r, column=i)
                if c.kind in FORMATS:
                    cell.number_format = FORMATS[c.kind]
                if flag == "low":
                    cell.font = LOW_FONT
                elif flag == "total":
                    cell.font = Font(bold=True)
                v = row.get(c.key)
                widths[i - 1] = max(widths[i - 1], min(len(str(v)) + 2 if v is not None else 0, 60))
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(w, 10)
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
        if not sec.rows:
            ws.append([t("reports.no_data")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
